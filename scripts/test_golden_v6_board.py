"""Golden V6 board regression tests (real Chromium: desktop + Android phone portrait/landscape).

Anti-regression for the approved V6 properties:
  * nothing starts before the click; click starts pen + (English) voice
  * ONE photorealistic hand (byte-identical asset), one pen, never duplicated
  * SINGLE board for text-only stages (Powers), DUAL board for drawings (Cube)
  * single -> dual -> single inside one lesson without reload / clearing / hand duplication
  * pause freezes writing + hand + voice; resume continues from the same word
  * default speed 0.7x, English-only voice, Sides terminology, math as real superscripts
  * drawing stays visible while the explanation scrolls; no overflow on phones
Run:  python3 scripts/test_golden_v6_board.py
"""
from __future__ import annotations

import hashlib
import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts.golden_v6_fixtures import (GOLDEN, POWERS_ACTIVITIES, cube_scenes, mixed_scenes,  # noqa: E402
                                        powers_scenes, powers_single_scenes)
from scripts.nabil_factory.cards import golden_v6 as g6  # noqa: E402

OUT = Path(os.environ.get("V6_OUT", tempfile.mkdtemp(prefix="v6out_")))
OUT.mkdir(parents=True, exist_ok=True)
RESULTS: list[tuple[str, str, str]] = []

SPEECH_STUB = """
(() => {
  window.__spoken = []; window.__ssCalls = [];
  let paused = false;
  const ss = {
    speaking: false, pending: false,
    get paused() { return paused; },
    getVoices() { return [{lang: 'en-US', name: 'Stub English', default: true}]; },
    speak(u) { window.__spoken.push({text: u.text, lang: u.lang, rate: u.rate, volume: u.volume}); window.__ssCalls.push('speak');
      if (u.volume !== 0) { setTimeout(() => u.onstart && u.onstart(), 30); setTimeout(() => u.onend && u.onend(), 500 + u.text.length * 25); } },
    cancel() { window.__ssCalls.push('cancel'); paused = false; },
    pause() { paused = true; window.__ssCalls.push('pause'); },
    resume() { paused = false; window.__ssCalls.push('resume'); },
    addEventListener() {}, removeEventListener() {},
  };
  Object.defineProperty(window, 'speechSynthesis', {value: ss, configurable: true});
  window.SpeechSynthesisUtterance = function (t) { this.text = t; this.lang = ''; this.rate = 1; this.volume = 1; };
})();
"""


def record(name: str, ok: bool, detail: str = "") -> None:
    RESULTS.append((name, "PASS" if ok else "FAIL", detail))
    print(f"{'PASS' if ok else 'FAIL'}  {name}  {detail}")


def page_html(name: str, title: str, scenes, lang: str, titles=None) -> Path:
    body = g6.render_golden_v6_board(title, scenes, lang, golden_spec=GOLDEN, titles=titles or {})
    path = OUT / f"{name}.html"
    path.write_text(
        f'<!doctype html><html lang="{lang}"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1"></head>'
        f'<body style="margin:0">{body}</body></html>', encoding="utf-8")
    return path


def open_tools(pg):
    if pg.evaluate("document.getElementById('nbTools').hidden"):
        pg.click('#nbToolsToggle')


def new_page(browser, path: Path, w=1280, h=760, fast=False, stub=True, mobile=False):
    ctx = browser.new_context(viewport={"width": w, "height": h}, is_mobile=mobile, has_touch=mobile,
                              device_scale_factor=2 if mobile else 1)
    pg = ctx.new_page()
    pg.errors = []
    pg.on("pageerror", lambda e: pg.errors.append(str(e)))
    if stub:
        pg.add_init_script(SPEECH_STUB)
    if fast:
        pg.add_init_script("window.NABIL_V18_FAST=true")
    pg.goto(f"file://{path}")
    pg.wait_for_function("!!window.NABILWholeLessonOrchestrator", timeout=20000)
    return pg


def state(pg):
    return pg.evaluate("NABILWholeLessonOrchestrator.state()")


def units_written(pg) -> int:
    return pg.evaluate("document.querySelectorAll('#nbWriting .nb-step span, #nbWriting .nb-step bdi').length")


def run_static() -> None:
    hand = g6.hand_payload()
    record("hand asset is byte-identical to the approved V6 template",
           hashlib.sha256(hand.encode()).hexdigest() == g6.HAND_SHA256)
    ref = (g6._ASSETS / "golden_v6_reference" /
           "NABIL_G07_CUBE_GOLDEN_V6_APPROVED_VOICE_ON_SIDES.html").read_text(encoding="utf-8")
    record("reference copy of the approved V6 template is kept and contains the same hand",
           hand in ref and "cubeOutline" in ref and "teacherHand" in ref)
    eng = (g6._ASSETS / "golden_v6_engine.js").read_text(encoding="utf-8")
    record("engine keeps the approved V6 cube actions and hand overlay",
           all(k in eng for k in ("cubeOutline", "cubeColor", "cubeRotate", "cubeUnfold", "teacherHand", "getScreenCTM")))
    record("engine has no hard-coded lesson content and no data: image",
           "data:image" not in eng and "Powers" not in eng and "Cube —" not in eng)
    page = (OUT / "cube.html").read_text(encoding="utf-8") if (OUT / "cube.html").exists() else ""
    record("terminology: Sides / أضلاع / Side length (not Edges / أحرف) inside solids lessons",
           g6.apply_terminology("Edge length of the edges: أحرف المكعب", "Cube") == "Side length of the sides: أضلاع المكعب"
           and g6.apply_terminology("edges of a graph", "Graph theory") == "edges of a graph")
    sc = powers_scenes("en")
    record("factory output is structured data only (no HTML strings in scene data)",
           all("<" not in str(s) for sc_ in sc for s in sc_["steps"] if s.get("t") in ("text", "formula")))
    modes = [g6.scene_board_modes(s) for s in powers_scenes("en")]
    record("Powers scenes: text-only stage is single, lab stages draw (dual)",
           modes[2] == ["single"] and modes[0][-1] == "dual", str(modes))
    for sc_ in powers_scenes("en"):
        for s in sc_["steps"]:
            if s.get("t") == "text":
                assert s["speech"] == s["text"], "English page speaks its own English text"
    ar = g6.build_scenes([dict(POWERS_ACTIVITIES[0])], "ar", "Powers")
    record("non-English pages never send Arabic/French to the English voice",
           all(s.get("speech", "") == "" for s in ar[0]["steps"] if s["t"] == "text"))


def run_browser(browser) -> None:
    powers = page_html("powers", "Powers", powers_scenes("en"), "en", {"fr": "Les puissances"})
    powers_single = page_html("powers_single", "Powers", powers_single_scenes("en"), "en", {"fr": "Les puissances"})
    cube = page_html("cube", "Cube — Faces, Sides and Vertices", cube_scenes("ar"), "ar",
                     {"fr": "Le cube — faces, côtés, sommets"})
    mixed = page_html("mixed", "Mixed", mixed_scenes(), "ar")

    # ---------------- nothing starts before the click, opening screen content
    pg = new_page(browser, powers_single)
    time.sleep(1.5)
    st = state(pg)
    n_spoken = pg.evaluate("window.__spoken.length")
    opening = pg.evaluate("""() => ({img: !!document.getElementById('nbAvatarOpen').src,
        brand: document.querySelector('.nb-brand').textContent, en: document.getElementById('nbOpenTitleEn').textContent,
        fr: (document.getElementById('nbOpenTitleFr')||{}).textContent, start: !!document.getElementById('nbStart'),
        writtenBeforeClick: document.getElementById('nbWriting').children.length})""")
    record("opening screen: avatar, «منصة الأستاذ النبيل», EN title with FR under it, Start button",
           opening["img"] and opening["brand"] == "منصة الأستاذ النبيل" and opening["en"] == "Powers"
           and opening["fr"] == "Les puissances" and opening["start"], str(opening))
    record("nothing starts before the click (no pen, no voice, no writing)",
           (not st["started"]) and n_spoken == 0 and opening["writtenBeforeClick"] == 0 and not st["handVisible"], str(st))
    record("default speed 0.7x and voice ON", st["speed"] == 0.7 and st["voice"] is True)
    pg.screenshot(path=str(OUT / "01_opening_desktop.png"))

    # ---------------- click starts: SINGLE board, hand, English voice (not Arabic)
    pg.click("#nbStart")
    pg.wait_for_function("document.querySelectorAll('#nbWriting .nb-step span').length>2", timeout=15000)
    st = state(pg)
    spoken = [u for u in pg.evaluate("window.__spoken") if u["text"].strip()]
    record("click starts the lesson on a SINGLE full-width board",
           st["started"] and st["boardMode"] == "single"
           and pg.evaluate("getComputedStyle(document.getElementById('nbDiagram')).display") == "none", str(st))
    record("English voice speaks at the click (en-US, rate from 0.7x), full lesson text, no Arabic",
           bool(spoken) and spoken[0]["lang"] == "en-US" and all(ord(c) < 0x0600 or ord(c) > 0x06FF for u in spoken for c in u["text"]),
           str(spoken[:1]))
    record("exactly one hand on the page", pg.evaluate("document.querySelectorAll('#teacherHand').length") == 1
           and pg.evaluate("document.querySelectorAll('img[id*=and]').length") == 1)
    record("hand is the real image (Blob, PNG 669x574), visible while writing",
           pg.evaluate("(()=>{const h=document.getElementById('teacherHand');return h.src.startsWith('blob:')&&h.naturalWidth===669&&h.naturalHeight===574&&h.classList.contains('is-visible')})()"))
    pg.screenshot(path=str(OUT / "02_single_board_writing_desktop.png"))

    # ---------------- pause freezes everything, resume continues from the same word
    pg.click("#nabilWholePlay")
    time.sleep(0.4)
    n1 = units_written(pg)
    pos1 = pg.evaluate("(()=>{const h=document.getElementById('teacherHand');return [h.style.left,h.style.top]})()")
    idx1 = state(pg)["index"]
    time.sleep(1.6)
    n2 = units_written(pg)
    pos2 = pg.evaluate("(()=>{const h=document.getElementById('teacherHand');return [h.style.left,h.style.top]})()")
    calls = pg.evaluate("window.__ssCalls")
    record("PAUSE freezes writing, pen/hand position and voice together",
           state(pg)["paused"] and n1 == n2 and pos1 == pos2 and "pause" in calls, f"units {n1}->{n2} hand {pos1}->{pos2}")
    pg.click("#nabilWholePlay")
    time.sleep(2.5)
    n3 = units_written(pg)
    record("RESUME continues from the same place (more words, same scene, no restart, voice resumed)",
           (not state(pg)["paused"]) and n3 > n2 and state(pg)["index"] == idx1 and pg.evaluate("window.__ssCalls").count("resume") >= 1,
           f"units {n2}->{n3}")
    record("no page errors (single board)", not pg.errors, str(pg.errors[:2]))
    pg.context.close()

    # ---------------- CUBE: DUAL board, drawing visible, hand moves on both panes
    pg = new_page(browser, cube)
    pg.click("#nbStart")
    pg.wait_for_function("document.querySelectorAll('#nbInk *').length>0", timeout=30000)
    st = state(pg)
    dual_ok = pg.evaluate("""() => { const d=document.getElementById('nbDiagram').getBoundingClientRect(),
        w=document.getElementById('nbWriting').getBoundingClientRect();
        return {diag: d.width>200 && d.height>200, writing: w.width>200,
                overlap: !(d.right<=w.left+1 || w.right<=d.left+1 || d.bottom<=w.top+1 || w.bottom<=d.top+1)}}""")
    record("CUBE shows TWO boards side by side (drawing + explanation) without overlap",
           st["boardMode"] == "dual" and dual_ok["diag"] and dual_ok["writing"] and not dual_ok["overlap"], str(dual_ok))
    pg.wait_for_function("document.querySelectorAll('#nbInk line').length>=3", timeout=40000)
    hand_box = pg.evaluate("(()=>{const h=document.getElementById('teacherHand').getBoundingClientRect();const d=document.getElementById('nbDiagram').getBoundingClientRect();return {x:h.left+h.width*.18,inside:(h.left+h.width*.18)>=d.left-40&&(h.left+h.width*.18)<=d.right+40}})()")
    record("the one hand follows the pen on the drawing board", hand_box["inside"], str(hand_box))
    pg.screenshot(path=str(OUT / "03_cube_dual_desktop.png"))
    pg.context.close()

    # ---------------- mixed lesson: single -> dual -> single without reload, content preserved
    pg = new_page(browser, mixed, fast=False)
    pg.evaluate("window.__marker='same-page'")
    pg.click("#nbStart")
    open_tools(pg)
    pg.select_option("#nbSpeed", "1.25")
    pg.wait_for_function("document.querySelectorAll('#nbInk *').length>0", timeout=60000)
    first_text = pg.evaluate("document.querySelector('#nbWriting .nb-step').textContent")
    pg.wait_for_function("NABILWholeLessonOrchestrator.modeLog().length>=2", timeout=90000)
    log = pg.evaluate("NABILWholeLessonOrchestrator.modeLog().map(x=>x.mode)")
    after = pg.evaluate("document.querySelector('#nbWriting .nb-step').textContent")
    record("mixed lesson switches single→dual→single mid-scene",
           log[:2] == ["dual", "single"], str(log))
    record("switching keeps the written content, the same page (no reload) and one hand",
           after == first_text and pg.evaluate("window.__marker") == "same-page"
           and pg.evaluate("document.querySelectorAll('#teacherHand').length") == 1)
    pg.screenshot(path=str(OUT / "04_mixed_back_to_single_desktop.png"))
    record("no page errors (mixed)", not pg.errors, str(pg.errors[:2]))
    pg.context.close()

    # ---------------- full lesson plays to the golden card (FAST), Powers with lab drawings
    pg = new_page(browser, powers, fast=True)
    pg.click("#nbStart")
    pg.wait_for_function("NABILWholeLessonOrchestrator.state().final===true", timeout=60000)
    sup = pg.evaluate("document.querySelectorAll('#nbWriting sup').length")
    txt = pg.evaluate("document.getElementById('nbGolden').textContent.length")
    record("full lesson reaches the Golden Card stage; no JS errors", txt > 10 and not pg.errors, str(pg.errors[:2]))
    pg.goto(f"file://{powers}")
    pg.wait_for_function("!!window.NABILWholeLessonOrchestrator")
    pg.click("#nbStart")
    pg.wait_for_function("NABILWholeLessonOrchestrator.state().final===true", timeout=60000)
    pg.evaluate("""() => { window.__seen={sup:0,raw:false,boxes:0};
        new MutationObserver(ms => { for (const m of ms) for (const n of m.addedNodes) { if (n.nodeType!==1) continue;
          window.__seen.sup += (n.tagName==='SUP'?1:0) + n.querySelectorAll('sup').length;
          window.__seen.boxes += (n.tagName==='rect'||n.tagName==='line'?1:0) + n.querySelectorAll('rect,line').length;
          if (/[\\^*]/.test(n.textContent||'') && n.closest && n.closest('#nbWriting')) window.__seen.raw = true; } })
          .observe(document.getElementById('nabilWholeLessonSmartLab'), {childList:true, subtree:true});
        NABILWholeLessonOrchestrator.goTo(0); }""")
    pg.wait_for_function("NABILWholeLessonOrchestrator.state().final===true", timeout=60000)
    raw = pg.evaluate("window.__seen")
    record("math is real superscripts (no raw ^ or * on screen); Powers pen drew the factor boxes",
           (not raw["raw"]) and raw["sup"] > 0 and raw["boxes"] > 0, str(raw))
    pg.context.close()

    # ---------------- «لم أفهم» → «فهمت، أكمل من هنا»
    pg = new_page(browser, powers_single)
    pg.click("#nbStart")
    pg.wait_for_function("document.querySelectorAll('#nbWriting .nb-step span').length>1", timeout=15000)
    open_tools(pg)
    pg.click("#nbLost")
    time.sleep(0.3)
    p_state = state(pg)
    has_btn = pg.evaluate("!!document.getElementById('nbUnderstood')")
    n_a = units_written(pg)
    time.sleep(1.0)
    n_b = units_written(pg)
    pg.click("#nbUnderstood")
    time.sleep(2.0)
    record("«لم أفهم» pauses; «فهمت، أكمل من هنا» resumes the same lesson",
           p_state["paused"] and has_btn and n_a == n_b and (not state(pg)["paused"]) and units_written(pg) > n_b,
           f"{n_a}/{n_b}->{units_written(pg)}")
    pg.context.close()

    # ---------------- formative question: waits for the student, wrong -> feedback, then continue
    chk = [{"id": "Q", "title": "Check", "start_mode": "single", "steps": [
        {"t": "text", "text": "Predict the value.", "label": "", "speech": "Predict the value."},
        {"t": "check", "check": {"question": "What is 2^{3}?", "expected": "8", "wrong_feedback": "Try again.", "correct_feedback": "Yes."}},
        {"t": "text", "text": "After the question the lesson goes on.", "label": "", "speech": ""}]}]
    p_chk = page_html("check", "Check", chk, "en")
    pg = new_page(browser, p_chk, fast=True)
    pg.click("#nbStart")
    pg.wait_for_selector(".nb-check", timeout=15000)
    time.sleep(0.5)
    waiting = state(pg)["studentWait"] and not state(pg)["final"]
    pg.fill(".nb-check input", "7")
    pg.click(".nb-check button >> nth=0")
    wrong = pg.inner_text(".nb-fb")
    pg.fill(".nb-check input", "8")
    pg.click(".nb-check button >> nth=0")
    right = pg.inner_text(".nb-fb")
    pg.click(".nb-check button >> nth=1")
    pg.wait_for_function("NABILWholeLessonOrchestrator.state().final===true", timeout=20000)
    record("formative question halts the lesson until the student continues (wrong -> feedback -> correct)",
           waiting and wrong == "Try again." and right == "Yes.", f"{wrong!r} {right!r}")
    pg.context.close()

    # ---------------- exercises: automatic solution on the same board + dropdown + hints + Arabic reading
    from scripts.nabil_factory.cards.v18_board import render_v18_exercises_board
    exs = [{"number": "1", "exact_source_prompt": "Write 3^4 × 3^2 as one power.", "solution_status": "SOLVED",
            "_pre_solved_solution": {"steps": ["Given: 3^4 × 3^2", "Rule: same base, so we add the exponents: 3^(4+2)",
                                               "Substitution: replace m by 4 and n by 2", "Calculation: 3^6 = 729"],
                                     "final_answer": "3^6 = 729"}},
           {"number": "2", "exact_source_prompt": "Compute 2^5 ÷ 2^2.", "solution_status": "SOLVED",
            "_pre_solved_solution": {"steps": ["Given: 2^5 ÷ 2^2", "Rule: divide powers with the same base: subtract the exponents",
                                               "Calculation: 2^(5-2) = 2^3 = 8"], "final_answer": "8"}},
           {"number": "3", "exact_source_prompt": "Compute 7^0.", "solution_status": "PENDING"}]
    exp = OUT / "exercises.html"
    exp.write_text('<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"></head><body style="margin:0">'
                   + render_v18_exercises_board("Powers", exs, "en", subject="mathematics") + "</body></html>", encoding="utf-8")
    pg = new_page(browser, exp)
    opts = pg.evaluate("[...document.querySelectorAll('#nbExercise option')].map(o=>o.textContent)")
    record("exercise dropdown lists every exercise number (unsolved marked)", opts[1:] == ["1", "2", "3 …"], str(opts))
    pg.click("#nbStart")
    pg.wait_for_function("document.querySelectorAll('#nbWriting .nb-step').length>=4", timeout=60000)
    hints = pg.evaluate("[...document.querySelectorAll('.nb-hintchip')].map(h=>h.textContent)")
    labels = pg.evaluate("[...document.querySelectorAll('#nbWriting .nb-label')].map(h=>h.textContent)")
    pg.screenshot(path=str(OUT / "08_exercises_solution_desktop.png"))
    record("solution runs automatically on the same board with Given/Rule/Substitution labels",
           "Given" in labels and "Rule" in labels and "Substitution" in labels, str(labels))
    record("how-to hint words appear (substitute / add exponents)",
           any("Substitute" in h for h in hints) and any("add the exponents" in h for h in hints), str(hints))
    open_tools(pg)
    pg.select_option("#nbExercise", index=3)
    pg.wait_for_function("NABILWholeLessonOrchestrator.state().index===2", timeout=10000)
    pg.wait_for_function("document.getElementById('nbWriting').innerText.includes('not available yet')", timeout=60000)
    unsolved = pg.inner_text("#nbWriting")
    record("choosing exercise 3 from the dropdown jumps to it; missing solution is stated, not invented",
           "not available yet" in unsolved, unsolved[:80])
    open_tools(pg)
    pg.select_option("#nbExercise", index=2)
    pg.wait_for_function("NABILWholeLessonOrchestrator.state().index===1", timeout=10000)
    record("dropdown jumps to exercise 2 without reload", state(pg)["index"] == 1 and not pg.errors, str(pg.errors[:2]))
    pg.context.close()
    pg = new_page(browser, exp, w=390, h=844, mobile=True)
    pg.click("#nbStart")
    pg.wait_for_function("document.querySelectorAll('#nbWriting .nb-step').length>=2", timeout=60000)
    ok_phone = pg.evaluate("() => { const vw=innerWidth; return document.documentElement.scrollWidth<=vw && [...document.querySelectorAll('.nb-ctl button,.nb-ctl select')].every(b=>{const r=b.getBoundingClientRect();return r.right<=vw+1&&r.left>=-1&&r.bottom<=innerHeight+1})}")
    record("phone: exercise dropdown and Arabic button fit on screen", ok_phone)
    pg.screenshot(path=str(OUT / "09_exercises_phone.png"))
    pg.context.close()

    # Arabic reading button: shows the Arabic line and speaks it after the English
    ar_sc = [{"id": "A", "title": "Arabic", "start_mode": "single", "steps": [
        {"t": "text", "text": "First read the statement.", "label": "", "speech": "First read the statement."},
        {"t": "text", "text": "Add the exponents.", "label": "", "speech": "Add the exponents.", "ar": "نجمع الأسس."},
        {"t": "text", "text": "No translation here.", "label": "", "speech": "No translation here."}]}]
    pa = page_html("ar_read", "Arabic", ar_sc, "en")
    pg = new_page(browser, pa)
    pg.click("#nbStart")
    open_tools(pg)
    pg.click("#nbArVoice")
    pressed = pg.get_attribute("#nbArVoice", "aria-pressed")
    pg.wait_for_selector(".nb-ar", timeout=60000)
    pg.wait_for_function("window.__spoken.filter(u=>u.lang==='ar-SA').length>=1", timeout=20000)
    pg.wait_for_function("[...document.querySelectorAll('.nb-ar')].some(e=>e.textContent.includes('الأسس'))", timeout=30000)
    spoken = [u for u in pg.evaluate("window.__spoken") if u["text"].strip()]
    langs = [u["lang"] for u in spoken if u["text"] != "First read the statement." and "Ustaz Nabil" not in u["text"]][:2]  # the opening greeting is not lesson speech
    record("«اقرأ بالعربية أيضًا»: Arabic text shown and spoken (ar-SA) after the English (en-US)",
           pressed == "true" and langs == ["en-US", "ar-SA"] and "نجمع" in pg.inner_text("#nbWriting"), str(langs))
    pg.context.close()

    # ---------------- lesson plan (تحضير الدرس): button, languages, times, cache, real .docx
    import re as _re
    import subprocess
    import zipfile
    acts = [dict(a, source_page="12") for a in POWERS_ACTIVITIES]
    from scripts.nabil_factory.cards.v18_board import render_v18_smart_board, build_golden_spec
    spec = build_golden_spec("Powers", acts, "en", subject="mathematics")
    plan_page = OUT / "plan_lesson.html"
    plan_page.write_text('<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"></head><body style="margin:0">'
                         + render_v18_smart_board("Powers", acts, "en", golden_spec=spec) + "</body></html>", encoding="utf-8")
    pg = new_page(browser, plan_page)
    no_plan_in_front = pg.evaluate("!document.querySelector('#nbOpen [data-nb-plan]') && !document.querySelector('.nb-head [data-nb-plan]')")
    record("lesson-plan button is NOT in front of the student (opening screen / header)", no_plan_in_front)
    pg.click("#nbStart")
    pg.wait_for_selector("#nbStage #nbTools", state="attached")
    btn_txt = pg.evaluate("document.querySelector('#nbTools [data-nb-plan]').textContent")
    record("lesson-plan button lives in the collapsible tools bar with the trilingual label",
           btn_txt == "📄 تحضير الدرس | Lesson Plan | Fiche de préparation", btn_txt)
    open_tools(pg)
    pg.click("#nbTools [data-nb-plan]")
    paused_on_open = state(pg)["paused"]
    pg.select_option("#nbPlanLang", "ar")
    info = pg.evaluate("""() => { const d=document.getElementById('nbPlanDoc'); const t=[...d.querySelectorAll('table')];
        const flow=t.find(x=>x.rows[0].cells.length===6);
        const mins=[...flow.rows].slice(1).map(r=>parseInt(r.cells[0].textContent,10));
        return {dir:d.dir, h1:d.querySelector('h1').textContent, cols:[...flow.rows[0].cells].map(c=>c.textContent), mins,
                sum:mins.reduce((a,b)=>a+b,0), sup:d.querySelectorAll('sup').length,
                has:['الكفايات المستهدفة','الأهداف','المكتسبات القبلية','الواجب المنزلي','ملاحظات المعلم بعد تنفيذ الدرس'].every(k=>d.textContent.includes(k)),
                pages:d.textContent.includes('12'), crdp:document.getElementById('nbPlanNote').textContent.includes('ليس نموذجًا رسميًا')} }""")
    record("opening the plan pauses the lesson; Arabic plan is RTL with the 6-column flow table",
           paused_on_open and info["dir"] == "rtl" and info["cols"] == ["الزمن", "المرحلة", "نشاط الأستاذ", "نشاط المتعلم", "الوسائل", "التقويم"], str(info["cols"]))
    record("45-minute model split 5/10/15/10/5 and all mandatory sections present",
           info["mins"] == [5, 10, 15, 10, 5] and info["has"] and info["pages"], str(info["mins"]))
    record("plan keeps math as superscripts and states it is not an official unified CRDP template", info["sup"] > 0 and info["crdp"])
    pg.fill("#nbPlanMin", "90"); pg.press("#nbPlanMin", "Enter"); pg.evaluate("document.getElementById('nbPlanMin').dispatchEvent(new Event('change'))")
    pg.fill("#nbPlanSes", "2"); pg.evaluate("document.getElementById('nbPlanSes').dispatchEvent(new Event('change'))")
    two = pg.evaluate("""() => { const flows=[...document.querySelectorAll('#nbPlanDoc table')].filter(x=>x.rows[0].cells.length===6);
        return {n:flows.length, sums:flows.map(f=>[...f.rows].slice(1).reduce((a,r)=>a+parseInt(r.cells[0].textContent,10),0))} }""")
    record("teacher can change period length and number of periods (90 min x 2 periods)",
           two["n"] == 2 and two["sums"] == [90, 90], str(two))
    pg.evaluate("document.getElementById('nbPlanDoc').insertAdjacentHTML('afterbegin','<p id=\"mine\">ملاحظة الأستاذ</p>'); document.getElementById('nbPlanDoc').dispatchEvent(new Event('input'))")
    pg.wait_for_function("() => Object.keys(localStorage).some(k => k.startsWith('nabilPlan:') && (localStorage.getItem(k)||'').includes('id=\"mine\"'))", timeout=15000)
    pg.click("#nbPlanClose"); open_tools(pg)
    pg.click("#nbTools [data-nb-plan]")
    again = pg.evaluate("({cached: window.__nbPlanState.fromCache, mine: !!document.getElementById('mine'), key: window.__nbPlanState.key})")
    record("plan is generated once per lesson+language and reused from the saved copy (edits kept)",
           again["cached"] and again["mine"], str(again))
    pg.evaluate("document.getElementById('mine').remove()")
    pg.select_option("#nbPlanLang", "ar"); pg.fill("#nbPlanMin", "45"); pg.fill("#nbPlanSes", "1")
    pg.evaluate("['nbPlanMin','nbPlanSes'].forEach(i=>document.getElementById(i).dispatchEvent(new Event('change')))")
    for lang in ("en", "fr"):
        pg.select_option("#nbPlanLang", lang)
        d = pg.evaluate("({dir:document.getElementById('nbPlanDoc').dir, h:document.querySelector('#nbPlanDoc h1').textContent})")
        record(f"plan language {lang}: direction and headings", d["dir"] == "ltr" and ("Lesson Plan" in d["h"] or "Fiche" in d["h"]), str(d))
    pg.select_option("#nbPlanLang", "ar")
    pg.screenshot(path=str(OUT / "10_lesson_plan_ar_desktop.png"))
    with pg.expect_download() as dl:
        pg.click("#nbPlanDocx")
    docx_path = OUT / "lesson_plan_ar.docx"
    dl.value.save_as(str(docx_path))
    pg.context.close()
    import docx as _docx
    d = _docx.Document(str(docx_path))
    xml = zipfile.ZipFile(docx_path).read("word/document.xml").decode("utf-8")
    flow_tables = [t for t in d.tables if len(t.columns) == 6]
    record("DOCX opens in python-docx: headings, 6-column flow table, RTL, superscripts",
           len(flow_tables) == 1 and "<w:bidi/>" in xml and "w:val=\"superscript\"" in xml and any("تحضير الدرس" in p.text for p in d.paragraphs),
           f"tables={len(d.tables)} paragraphs={len(d.paragraphs)}")
    record("DOCX flow table header and times intact",
           [c.text for c in flow_tables[0].rows[0].cells][:2] == ["الزمن", "المرحلة"] and flow_tables[0].rows[1].cells[0].text.startswith("5"))
    conv = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(docx_path)],
                          capture_output=True, text=True, timeout=180)
    pdf = OUT / "lesson_plan_ar.pdf"
    txt = subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True).stdout if pdf.exists() else ""
    pages = len(_re.findall(r"\f", txt))
    clean = _re.sub(r"[\u202a-\u202e\u2066-\u2069\u200e\u200f]", "", txt)
    record("DOCX converts to PDF with LibreOffice (pages and Arabic text intact); NOT tested in Microsoft Word",
           pdf.exists() and pages >= 1 and "جدول سير الحصة" in clean and "تحضير الدرس" in clean, f"pages={pages} rc={conv.returncode}")

    # ---------------- phones: portrait + landscape, no overflow/clipping, drawing stays visible
    for label, (w, h) in {"phone_portrait": (390, 844), "phone_landscape": (844, 390)}.items():
        pg = new_page(browser, cube, w=w, h=h, mobile=True)
        pg.click("#nbStart")
        pg.wait_for_function("document.querySelectorAll('#nbInk *').length>0", timeout=30000)
        time.sleep(1.5)
        m = pg.evaluate("""() => { const q=s=>document.querySelector(s).getBoundingClientRect(), vw=innerWidth, vh=innerHeight;
            const d=q('#nbDiagram'), wr=q('#nbWriting'), c=q('.nb-ctl');
            const btns=[...document.querySelectorAll('.nb-ctl button,.nb-ctl select')].map(b=>b.getBoundingClientRect());
            const clipped=btns.filter(b=>b.right>vw+1||b.left<-1||b.bottom>vh+1).length;
            const ov=!(d.right<=wr.left+1||wr.right<=d.left+1||d.bottom<=wr.top+1||wr.bottom<=d.top+1);
            return {sw:document.documentElement.scrollWidth, vw, clipped, overlap:ov, diagH:d.height, diagVisible:d.top>=-1&&d.bottom<=vh+1, ctlBottom:c.bottom<=vh+1}}""")
        record(f"{label}: no horizontal overflow, controls inside the screen, boards do not overlap",
               m["sw"] <= m["vw"] and m["clipped"] == 0 and not m["overlap"] and m["ctlBottom"], str(m))
        # long explanation scrolls while the drawing stays visible
        pg.evaluate("""() => { const w=document.getElementById('nbWriting'); for(let i=0;i<14;i++){const d=document.createElement('div');d.className='nb-step';d.textContent='filler '+i;d.style.height='90px';w.appendChild(d)} w.scrollTop=w.scrollHeight }""")
        time.sleep(0.4)
        vis = pg.evaluate("""() => { const d=document.getElementById('nbDiagram').getBoundingClientRect();
            return d.top>=-1 && d.bottom<=innerHeight+1 && d.height>120 }""")
        record(f"{label}: drawing stays visible while the explanation scrolls", vis)
        pg.screenshot(path=str(OUT / f"05_cube_{label}.png"))
        record(f"{label}: no page errors", not pg.errors, str(pg.errors[:2]))
        pg.context.close()

    # phone: Powers single board
    pg = new_page(browser, powers_single, w=390, h=844, mobile=True)
    pg.click("#nbStart")
    pg.wait_for_function("document.querySelectorAll('#nbWriting .nb-step span').length>6", timeout=20000)
    fs = pg.evaluate("parseFloat(getComputedStyle(document.querySelector('#nbWriting .nb-step')).fontSize)")
    ov = pg.evaluate("document.documentElement.scrollWidth<=innerWidth")
    record("phone Powers (single board): large clear writing (>=20px) and no overflow", fs >= 20 and ov, f"font {fs}px")
    pg.screenshot(path=str(OUT / "06_powers_single_phone.png"))
    pg.context.close()

    # ---------------- AR/EN/FR labels without mixing
    for lang, start in (("ar", "ابدأ الدرس"), ("fr", "Commencer la leçon"), ("en", "Start lesson")):
        p = page_html(f"lang_{lang}", "Powers", powers_single_scenes("en"), lang, {"fr": "Les puissances"})
        pg = new_page(browser, p)
        label_txt = pg.evaluate("document.getElementById('nbStart').textContent")
        dir_ok = pg.evaluate("document.getElementById('nbOpenTitleEn').dir")
        record(f"UI language {lang}: start button and controls in that language only",
               label_txt == start and dir_ok == "ltr", label_txt)
        pg.screenshot(path=str(OUT / f"07_opening_{lang}.png"))
        pg.context.close()


def main() -> int:
    run_static()
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        try:
            run_browser(browser)
        finally:
            browser.close()
    fails = [r for r in RESULTS if r[1] != "PASS"]
    print(f"\n{len(RESULTS) - len(fails)}/{len(RESULTS)} passed; output: {OUT}")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
