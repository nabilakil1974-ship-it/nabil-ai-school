"""Upload an existing failed-pilot draft for human inspection only.

This never calls the lesson generator, changes a production ledger, or marks a
draft scientifically verified. It requires an existing saved draft JSON.
"""
import argparse
import base64
import html
import io
import json
import tempfile
from pathlib import Path

from scripts import nabil_lesson_factory as factory

BOOK_ID = "1LasqIgGUuck1l-2EZbj2kA0Dg9ygJ_AH"
WARNING = "UNVERIFIED PILOT DRAFT — Scientific Review failed. Answers may be wrong."


def find_draft(report_path):
    report = json.loads(Path(report_path).read_text(encoding="utf-8"))
    for attempt in report.get("attempts", []):
        if attempt.get("book_id") != BOOK_ID:
            continue
        for candidate in attempt.get("draft_paths", []):
            path = Path(candidate)
            if path.is_file() and path.parent == Path("/tmp"):
                return path
    raise FileNotFoundError("NO_SAVED_PILOT_DRAFT: nothing can be uploaded")


def preview_html(lesson, images):
    e = lambda value: html.escape(str(value or ""), quote=True)
    blocks = [f'<header><strong>{e(WARNING)}</strong></header><main>',
              f'<h1>{e(lesson.get("title", "Solids and Liquids"))}</h1>',
              f'<p>{e(lesson.get("introduction", ""))}</p>']
    for field, heading in (("concepts", "Concepts"), ("activities", "Activities"),
                           ("questions", "Questions"), ("exercises", "Exercises"),
                           ("summary", "Summary")):
        blocks.append(f'<h2>{heading}</h2>')
        for item in lesson.get(field, []):
            if not isinstance(item, dict):
                continue
            blocks.append('<section>')
            for key in ("heading", "prompt", "explanation", "answer", "solution", "text"):
                if item.get(key):
                    blocks.append(f'<p><b>{key}:</b> {e(item[key])}</p>')
            for key in ("options", "solution_steps"):
                if isinstance(item.get(key), list):
                    blocks.append('<ol>' + ''.join(f'<li>{e(value)}</li>' for value in item[key]) + '</ol>')
            if item.get("pdf_page"):
                blocks.append(f'<small>Draft citation: PDF page {e(item["pdf_page"])}</small>')
            blocks.append('</section>')
    blocks.append('<h2>Original textbook pages 13–18</h2>')
    for page, image in images.items():
        blocks.append(f'<figure><img alt="Original PDF page {page}" src="data:image/jpeg;base64,'
                      + base64.b64encode(image).decode("ascii") + f'"><figcaption>PDF page {page}</figcaption></figure>')
    blocks.append(f'<footer>{e(WARNING)}<br>Source: G 07 physics.pdf · {BOOK_ID}</footer></main>')
    css = ('body{margin:0;background:#081b31;color:#edf8ff;font:18px Arial,sans-serif;line-height:1.55}'
           'header{position:sticky;top:0;background:#921d29;padding:16px;z-index:2}'
           'main{max-width:960px;margin:auto;padding:18px}section,figure{background:#12314e;padding:16px;'
           'border-radius:14px;margin:16px 0}img{width:100%;height:auto}small{color:#a8dfef}'
           'h1,h2{color:#6debd1}footer{background:#921d29;padding:18px}')
    return ('<!doctype html><html lang="en"><head><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width,initial-scale=1">'
            '<title>UNVERIFIED — Solids and Liquids</title><style>' + css + '</style></head><body>'
            + '\n'.join(blocks) + '</body></html>').encode('utf-8')


def preview_pptx(lesson, images):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    def slide(title, body="", image=None):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        box = s.shapes.add_textbox(Inches(.5), Inches(.25), Inches(12.3), Inches(.5))
        box.text_frame.text = title
        box.text_frame.paragraphs[0].font.size = Pt(26)
        alert = s.shapes.add_textbox(Inches(.5), Inches(6.85), Inches(12.3), Inches(.4))
        alert.text_frame.text = WARNING
        alert.text_frame.paragraphs[0].font.size = Pt(13)
        if image:
            s.shapes.add_picture(io.BytesIO(image), Inches(2.2), Inches(.9), height=Inches(5.75))
        elif body:
            box = s.shapes.add_textbox(Inches(.7), Inches(1), Inches(11.9), Inches(5.6))
            box.text_frame.word_wrap = True
            box.text_frame.text = body[:1700]
            for p in box.text_frame.paragraphs:
                p.font.size = Pt(20)
    slide(lesson.get("title", "Solids and Liquids"), lesson.get("introduction", ""))
    for field in ("concepts", "activities", "exercises", "summary"):
        for item in lesson.get(field, []):
            if isinstance(item, dict):
                title = item.get("heading") or item.get("prompt") or field.title()
                body = '\n'.join(str(item.get(k, '')) for k in
                                 ("explanation", "answer", "solution", "text") if item.get(k))
                slide(str(title), body)
    for page, image in images.items():
        slide(f"Original textbook PDF page {page}", image=image)
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()


def main():
    from googleapiclient.http import MediaIoBaseUpload
    p = argparse.ArgumentParser()
    p.add_argument("--failed-report", default="/tmp/solids-liquids-dry-run.json")
    args = p.parse_args()
    draft = find_draft(args.failed_report)
    lesson = json.loads(draft.read_text(encoding="utf-8"))
    if lesson.get("title", "").strip().casefold() != "solids and liquids":
        raise ValueError("DRAFT_TITLE_MISMATCH")
    service = factory.owner_drive()
    with tempfile.TemporaryDirectory(prefix="nabil_unverified_preview_") as tmp:
        pdf = Path(tmp) / "source.pdf"
        factory.download_pdf_to_path(service, BOOK_ID, pdf)
        images = factory.source_images(pdf, [(p, "") for p in range(13, 19)])
        files = {
            "UNVERIFIED-Solids-and-Liquids.html": (preview_html(lesson, images), "text/html"),
            "UNVERIFIED-Solids-and-Liquids.pptx": (preview_pptx(lesson, images),
                                                   "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
        }
    folder = factory.ensure_folder(service, factory.ROOT_FOLDER, "NABIL_PILOT_TEST_UNVERIFIED")
    result = []
    for name, (raw, mime) in files.items():
        created = service.files().create(
            body={"name": name, "parents": [folder],
                  "description": WARNING + " Source PDF pages 13–18; production ledger unchanged."},
            media_body=MediaIoBaseUpload(io.BytesIO(raw), mimetype=mime, resumable=False),
            fields="id,name,size,mimeType,parents,webViewLink").execute()
        if int(created.get("size", 0)) != len(raw) or folder not in created.get("parents", []):
            raise RuntimeError("PREVIEW_UPLOAD_READBACK_FAILED")
        result.append({"name": name, "file_id": created["id"], "url": created.get("webViewLink")})
    print(json.dumps({"status": "UNVERIFIED_PREVIEW_UPLOADED", "source_draft": str(draft),
                      "folder_id": folder, "files": result}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
