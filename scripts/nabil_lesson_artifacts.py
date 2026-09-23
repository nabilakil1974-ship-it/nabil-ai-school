"""Render source-illustrated, bilingual HTML/PPTX artifacts for a review draft.

Every image is a faithful raster of a cited ORIGINAL PDF page, not a model-
invented scientific diagram. Source page images are reference material and
must be reviewed/cropped before final approval. No upload from this module.
"""
import io
import json
import re
from pathlib import Path

import fitz
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(7, 22, 34)
CARD = RGBColor(8, 36, 58)
CYAN = RGBColor(40, 200, 255)
WHITE = RGBColor(238, 249, 255)
GREEN = RGBColor(34, 229, 139)
YELLOW = RGBColor(255, 212, 59)
SLIDE_W = 13.333
SLIDE_H = 7.5

def render_source_pages(pdf_bytes, pages, out):
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    result = {}
    for page in sorted(set(pages)):
        if page < 1 or page > len(doc):
            raise ValueError("INVALID_PDF_PAGE")
        name = "source_pdf_page_%03d.png" % page
        pix = doc[page-1].get_pixmap(matrix=fitz.Matrix(1.3, 1.3),
                                     alpha=False)
        pix.save(str(out / name))
        result[page] = name
    return result

def add_source_figures(html, draft, pages):
    """Append an original source page directly under its related concept."""
    soup = BeautifulSoup(html, "html.parser")
    cards = soup.select("main section.card")
    concepts = draft["concepts"]
    if len(cards) < 1 + len(concepts):
        raise ValueError("LESSON_CONCEPT_CARDS_MISSING")
    for i, concept in enumerate(concepts):
        name = pages[concept["source_pdf_page"]]
        figure = soup.new_tag("figure")
        figure["class"] = "fig source-page"
        figure["data-source-pdf-page"] = str(concept["source_pdf_page"])
        img = soup.new_tag("img", src=name, loading="lazy")
        img["alt"] = "Original textbook PDF page %d" % concept["source_pdf_page"]
        figure.append(img)
        caption = soup.new_tag("figcaption")
        caption.string = "Original textbook · PDF p. %d · reference image, not a reconstructed diagram" % concept["source_pdf_page"]
        figure.append(caption)
        cards[i+1].append(figure)
    style = soup.new_tag("style")
    style.string = """.source-page{margin:15px 0;background:#08243a;border:2px solid #28c8ff;
border-radius:14px;padding:10px}.source-page img{display:block;width:100%;height:auto;
max-width:850px;margin:auto}.source-page figcaption{color:#ffd43b;font-size:.85rem}
@media(max-width:600px){.source-page{padding:5px}}"""
    soup.head.append(style)
    return str(soup)

def _rect(slide, x, y, w, h, color, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def _text(slide, text, x, y, w, h, size=22, color=WHITE, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(.12)
    tf.margin_right = Inches(.12)
    tf.margin_top = Inches(.07)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = str(text)
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box

def _slide(prs, title, body="", image=None, footer="", accent=CYAN):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    _rect(slide, 0, 0, .12, SLIDE_H, accent)
    _rect(slide, .35, .3, 12.55, .9, CARD, True)
    _text(slide, title, .55, .38, 12.0, .7, 27, accent, True)
    if image:
        _rect(slide, 6.55, 1.43, 6.35, 5.35, CARD, True)
        slide.shapes.add_picture(str(image), Inches(6.7), Inches(1.55),
                                 width=Inches(6.05), height=Inches(5.05))
        _rect(slide, .35, 1.43, 5.95, 5.35, CARD, True)
        _text(slide, body, .55, 1.67, 5.5, 4.75, 20)
    else:
        _rect(slide, .35, 1.43, 12.55, 5.35, CARD, True)
        _text(slide, body, .65, 1.7, 11.8, 4.7, 24)
    _text(slide, footer, .55, 7.0, 12.0, .3, 10, YELLOW)
    return slide

def make_pptx(draft, pages, out, lang, grade, subject, chapter):
    if lang not in ("fr","en"):
        raise ValueError("INVALID_LANGUAGE")
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    title = draft["title"][lang]
    base = "NABIL AI · Grade %s · %s · Chapter %s" % (grade, subject, chapter)
    _slide(prs, title, base, footer="Source-illustrated review draft · NOT APPROVED")
    objectives = "\n".join("• "+x[lang] for x in draft["objectives"])
    _slide(prs, "Objectifs" if lang=="fr" else "Objectives", objectives, footer=base)
    for x in draft["concepts"]:
        _slide(prs, x["heading"][lang], x["explanation"][lang],
               out/pages[x["source_pdf_page"]],
               "Original textbook · PDF page "+str(x["source_pdf_page"]))
    for i,x in enumerate(draft["activities"],1):
        _slide(prs, ("Activité " if lang=="fr" else "Activity ")+str(i),
               x["prompt"][lang], out/pages[x["source_pdf_page"]],
               "Original textbook · PDF page "+str(x["source_pdf_page"]))
        _slide(prs, ("Solution de l’activité " if lang=="fr" else "Activity solution ")+str(i),
               x["solution"][lang], footer=base, accent=GREEN)
    for x in draft["exercises"]:
        _slide(prs, ("Exercice " if lang=="fr" else "Exercise ")+str(x["number"]),
               x["prompt"][lang], out/pages[x["source_pdf_page"]],
               "Original textbook · PDF page "+str(x["source_pdf_page"]))
        _slide(prs, ("Solution de l’exercice " if lang=="fr" else "Exercise solution ")+str(x["number"]),
               x["solution"][lang], footer=base, accent=GREEN)
    for i,x in enumerate(draft["worksheet"],1):
        _slide(prs, ("Fiche de travail " if lang=="fr" else "Worksheet ")+str(i),
               x["question"][lang], footer=base)
        _slide(prs, ("Réponse " if lang=="fr" else "Answer ")+str(i),
               x["answer"][lang], footer=base, accent=GREEN)
    _slide(prs, "Carte de synthèse" if lang=="fr" else "Final reference card",
           "\n".join("• "+x[lang] for x in draft["summary"]),
           footer=base, accent=GREEN)
    path = out / ("lesson_"+lang+".pptx")
    prs.save(str(path))
    return path

def build_bundle(draft, pdf_bytes, out, html, grade, subject, chapter, title):
    citations = [x for k in ("concepts","activities","exercises")
                 for x in draft[k]]
    pages = render_source_pages(pdf_bytes,
              [x["source_pdf_page"] for x in citations],out)
    (out/"lesson.html").write_text(add_source_figures(html,draft,pages),
                                    encoding="utf-8")
    en = make_pptx(draft,pages,out,"en",grade,subject,chapter)
    fr = make_pptx(draft,pages,out,"fr",grade,subject,chapter)
    return {"html":"lesson.html","pptx_en":en.name,"pptx_fr":fr.name,
            "images":list(pages.values()),
            "source_claims":[{"pdf_page":x["source_pdf_page"],
                "verbatim_excerpt":x["verbatim_excerpt"],
                "lesson_claim":str(x.get("heading",x.get("prompt")))}
                for x in citations],
            "scientific_review_approved":False,
            "reference_design_approved":False}
