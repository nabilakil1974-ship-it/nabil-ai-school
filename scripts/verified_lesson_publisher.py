"""Fail-closed publishing gate for authored NABIL lesson bundles.

This is NOT a universal lesson generator. The source-grounded authoring stage
must supply a manifest and artifacts. Nothing is approved without source
excerpts, browser interaction checks, and an explicit --publish.
"""
import argparse
import hashlib
import io
import json
import re
from pathlib import Path

ROOT_ID = "16bcmZMO_dn4FqlGaDtl8Hky6iSBEqZpX"
SUBJECTS = {
    "mathematics": "Mathematics - رياضيات",
    "physics": "Physics",
    "chemistry": "Chemistry - كيمياء",
    "biology": "Biology - علوم الحياة",
    "general_science": "General Science - علوم عامة",
}
REQUIRED = ("lesson_id", "grade", "subject", "chapter", "title",
            "source_pdf", "source_sha256", "source_claims", "html",
            "pptx_en", "pptx_fr", "images")


def digest(path):
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def source_gate(manifest, base):
    from pypdf import PdfReader
    pdf = base / manifest["source_pdf"]
    if not pdf.is_file() or digest(pdf) != manifest["source_sha256"]:
        raise ValueError("SOURCE_PDF_MISSING_OR_HASH_MISMATCH")
    reader = PdfReader(str(pdf))
    claims = manifest["source_claims"]
    if not claims or not isinstance(claims, list):
        raise ValueError("SOURCE_CLAIMS_REQUIRED")
    for item in claims:
        page = item["pdf_page"]
        excerpt = item["verbatim_excerpt"]
        if not isinstance(page, int) or not 1 <= page <= len(reader.pages):
            raise ValueError("SOURCE_PAGE_INVALID")
        if len(excerpt.strip()) < 12:
            raise ValueError("SOURCE_EXCERPT_TOO_SHORT")
        actual = reader.pages[page - 1].extract_text() or ""
        normalize = lambda s: re.sub(r"\s+", " ", s).casefold().strip()
        if normalize(excerpt) not in normalize(actual):
            raise ValueError("SOURCE_EXCERPT_NOT_IN_PDF_PAGE_" + str(page))
        if not item.get("lesson_claim"):
            raise ValueError("SOURCE_CLAIM_MISSING")
    return {"pdf_sha256": digest(pdf), "claims_checked": len(claims)}


def artifact_gate(manifest, base):
    from bs4 import BeautifulSoup
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    html_path = base / manifest["html"]
    html = html_path.read_text(encoding="utf-8")
    soup = BeautifulSoup(html, "html.parser")
    if not soup.find("meta", attrs={"name": "viewport"}) or not soup.find("main"):
        raise ValueError("LESSON_STRUCTURE_MISSING")
    if not soup.select("canvas,svg,img") or not soup.select("input,button,select,details"):
        raise ValueError("FIGURE_OR_INTERACTIVITY_MISSING")
    if not soup.select("[data-en][data-fr]"):
        raise ValueError("BILINGUAL_CONTENT_MISSING")
    if not soup.select(".summary"):
        raise ValueError("FINAL_SUMMARY_MISSING")
    if not soup.select("details"):
        raise ValueError("WORKED_SOLUTION_MISSING")
    if not manifest["images"]:
        raise ValueError("NO_VERIFIED_FIGURES_OR_SOURCE_PAGE_IMAGES")
    for name in manifest["images"]:
        if not (base / name).is_file():
            raise ValueError("MISSING_IMAGE_" + name)
        if name not in html:
            raise ValueError("IMAGE_NOT_REFERENCED_" + name)
    pptx = {}
    for lang in ("en", "fr"):
        path = base / manifest["pptx_" + lang]
        prs = Presentation(str(path))
        if len(prs.slides) < 9:
            raise ValueError("PPTX_TOO_FEW_SLIDES_" + lang)
        if sum(any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in slide.shapes)
               for slide in prs.slides) < 4:
            raise ValueError("PPTX_INSUFFICIENT_VISUAL_SLIDES_" + lang)
        pptx[lang] = {"slides": len(prs.slides), "sha256": digest(path)}
    return {"html_sha256": digest(html_path), "pptx": pptx}


def browser_gate(manifest, base):
    from playwright.sync_api import sync_playwright
    url = (base / manifest["html"]).resolve().as_uri()
    screenshots = {}
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        try:
            for language in ("en", "fr"):
                for width, height in ((1280, 800), (375, 667)):
                    page = browser.new_page(viewport={"width": width, "height": height})
                    errors = []
                    page.on("pageerror", lambda error: errors.append(str(error)))
                    page.goto(url, wait_until="networkidle")
                    selector = page.locator("#lesson-language")
                    if selector.count() != 1:
                        raise ValueError("LANGUAGE_CONTROL_MISSING")
                    selector.select_option(language)
                    for details in page.locator("details").all():
                        details.locator("summary").click()
                    page.wait_for_timeout(150)
                    overflow = page.evaluate(
                        "() => document.documentElement.scrollWidth > document.documentElement.clientWidth + 1")
                    if overflow:
                        raise ValueError("HORIZONTAL_OVERFLOW_" + language + "_" + str(width))
                    if errors:
                        raise ValueError("BROWSER_JS_ERROR_" + repr(errors[:2]))
                    if page.locator("main").inner_text().strip() == "":
                        raise ValueError("EMPTY_RENDER")
                    screenshots[language + "_" + str(width)] = "passed"
                    page.close()
        finally:
            browser.close()
    return screenshots


def child(service, parent, name, mime=None):
    from scripts.nabil_lesson_factory import children
    matches = [f for f in children(service, parent)
               if f["name"].casefold() == name.casefold()
               and (mime is None or f["mimeType"] == mime)]
    if len(matches) != 1:
        raise ValueError("DRIVE_FOLDER_MISSING_OR_AMBIGUOUS_" + name)
    return matches[0]["id"]


def publish(manifest, base):
    from googleapiclient.http import MediaIoBaseUpload
    from scripts.nabil_lesson_factory import drive, children, MIME_FOLDER
    service = drive(write=True)
    grade = manifest["grade"]
    if not isinstance(grade, int) or not 1 <= grade <= 12:
        raise ValueError("GRADE_INVALID")
    subject = SUBJECTS[manifest["subject"]]
    grade_id = child(service, ROOT_ID, "Grade " + str(grade), MIME_FOLDER)
    subject_id = child(service, grade_id, subject, MIME_FOLDER)
    folder_name = f'{int(manifest["chapter"]):02d} - {manifest["title"]}'
    existing = [f for f in children(service, subject_id)
                if f["name"].casefold() == folder_name.casefold()]
    if existing:
        if len(existing) != 1 or existing[0]["mimeType"] != MIME_FOLDER:
            raise ValueError("DESTINATION_AMBIGUOUS")
        folder_id = existing[0]["id"]
    else:
        folder_id = service.files().create(
            body={"name": folder_name, "mimeType": MIME_FOLDER,
                  "parents": [subject_id]}, fields="id").execute()["id"]
    names = [manifest["html"], manifest["pptx_en"], manifest["pptx_fr"],
             *manifest["images"]]
    seen = {f["name"].casefold() for f in children(service, folder_id)}
    if any(Path(name).name.casefold() in seen for name in names):
        raise FileExistsError("REFUSE_OVERWRITE_EXISTING_DRIVE_ARTIFACTS")
    result = []
    for name in names:
        path = base / name
        mime = ("text/html" if name == manifest["html"] else
                "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                if name.endswith(".pptx") else "image/png")
        media = MediaIoBaseUpload(io.BytesIO(path.read_bytes()),
                                  mimetype=mime, resumable=False)
        uploaded = service.files().create(
            body={"name": path.name, "parents": [folder_id]},
            media_body=media, fields="id,name,size").execute()
        result.append(uploaded)
    return {"folder_id": folder_id, "files": result}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--manifest", required=True)
    ap.add_argument("--publish", action="store_true")
    ap.add_argument("--report", default="")
    args = ap.parse_args()
    path = Path(args.manifest).resolve()
    base = path.parent
    report = {"status": "REJECTED", "manifest": str(path)}
    try:
        manifest = json.loads(path.read_text(encoding="utf-8"))
        for key in REQUIRED:
            if key not in manifest:
                raise ValueError("MANIFEST_MISSING_" + key)
        report["source"] = source_gate(manifest, base)
        report["artifacts"] = artifact_gate(manifest, base)
        report["browser"] = browser_gate(manifest, base)
        report["status"] = "LOCALLY_VERIFIED_NOT_PUBLISHED"
        if args.publish:
            if manifest.get("scientific_review_approved") is not True or not manifest.get("reference_design_approved") is True:
                raise ValueError("SCIENTIFIC_AND_REFERENCE_DESIGN_APPROVAL_REQUIRED")
            report["drive"] = publish(manifest, base)
            report["status"] = "UPLOADED_PENDING_RAILWAY_ACCEPTANCE"
    except Exception as exc:
        report["error"] = type(exc).__name__ + ": " + str(exc)
    if args.report:
        Path(args.report).write_text(json.dumps(report, ensure_ascii=False, indent=2),
                                     encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] != "REJECTED" else 2


if __name__ == "__main__":
    raise SystemExit(main())
