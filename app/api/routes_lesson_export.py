"""Export existing lesson cards, not invented lesson content, as PPTX or printable reference."""
import io
import base64
import re
from html import escape
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse, HTMLResponse
from pydantic import BaseModel, Field
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from bs4 import BeautifulSoup
from app.api.routes_interactive_lessons import _resolve, _service, _download

router = APIRouter(prefix="/lesson-export", tags=["lesson-export"])

class Cards(BaseModel):
    title: str = Field(max_length=200)
    cards: list[str] = Field(min_length=1, max_length=65)
    source: str = Field(default="بطاقات الدرس المعروضة", max_length=200)
    images: list[list[str]] = Field(default_factory=list, max_length=65)

def _plain(text):
    return re.sub(r"\s+", " ", BeautifulSoup(str(text), "html.parser").get_text(" ", strip=True)).strip()[:5000]

def _from_drive(grade, subject, lesson, language):
    item = _resolve(grade, subject, lesson, language)
    html = _download(_service(), item["drive_file_id"]).decode("utf-8-sig")
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "button"]):
        tag.decompose()
    title = _plain((soup.find("h1") or soup.title or item["lesson"]).get_text(" ", strip=True) if soup.find("h1") or soup.title else item["lesson"])
    nodes = soup.select("main section, article, .card, .row")
    if not nodes:
        nodes = soup.select("h2, h3, p, li")
    cards = []
    for node in nodes:
        if node.find_parent(["section", "article"]) and node.name in ("section", "article"):
            continue
        value = _plain(node.get_text(" ", strip=True))
        if len(value) >= 20 and value not in cards:
            cards.append(value)
    if not cards:
        raise HTTPException(422, "الدرس لا يحتوي بطاقات نصية قابلة للاستخراج؛ لا يمكن إنشاء عرض أمين للمصدر.")
    return Cards(title=title, cards=cards[:65], source="الدرس المحضّر في Google Drive")

def _picture_bytes(value):
    """Only inline browser-rendered PNG/JPEG; never fetch arbitrary remote URLs."""
    if not isinstance(value, str) or not re.match(r"^data:image/(?:png|jpeg);base64,", value):
        return None
    try:
        raw = base64.b64decode(value.split(",", 1)[1], validate=True)
        if len(raw) > 3_000_000 or len(raw) < 30:
            return None
        from PIL import Image
        img = Image.open(io.BytesIO(raw))
        img.verify()
        return raw
    except Exception:
        return None


def _pptx(payload):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    def slide(title, body, figures=()):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(9, 30, 52)
        heading = s.shapes.add_textbox(Inches(.65), Inches(.4), Inches(12), Inches(.85))
        p = heading.text_frame.paragraphs[0]
        p.text = title[:140]
        p.font.size = Pt(27)
        p.font.bold = True
        p.font.color.rgb = RGBColor(110, 226, 232)
        good = [raw for value in figures[:4] if (raw := _picture_bytes(value))]
        body_width = 7.0 if good else 11.8
        box = s.shapes.add_textbox(Inches(.8), Inches(1.5), Inches(body_width), Inches(5.55))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = body[:3000]
        p.font.size = Pt(19 if len(body)<430 else 15 if len(body)<900 else 11)
        p.font.color.rgb = RGBColor(246, 250, 255)
        if good:
            # Preserve the source image aspect ratio, with no invented points or labels.
            from PIL import Image
            for i, raw in enumerate(good[:2]):
                try:
                    with Image.open(io.BytesIO(raw)) as im:
                        w, h = im.size
                    max_w, max_h = 4.55, 2.55 if len(good)>1 else 5.25
                    scale = min(max_w/w, max_h/h)
                    width, height = w*scale, h*scale
                    s.shapes.add_picture(io.BytesIO(raw), Inches(8.45+(max_w-width)/2),
                        Inches(1.5+i*2.8), width=Inches(width), height=Inches(height))
                except Exception:
                    continue
    slide(payload.title, payload.source)
    for i, card in enumerate(payload.cards, 1):
        slide(f"{payload.title} · {i}", _plain(card),
              payload.images[i-1] if i-1 < len(payload.images) else ())
    slide("البطاقة المرجعية | Révision", "\\n\\n".join(_plain(x)[:300] for x in payload.cards[-5:]))
    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return StreamingResponse(out,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": 'attachment; filename="NABIL_Lesson.pptx"',
                 "Cache-Control": "no-store"})

def _reference(payload):
    blocks="".join("<section><h2>بطاقة "+str(i)+"</h2><p>"+escape(_plain(card))+"</p></section>" for i,card in enumerate(payload.cards,1))
    return HTMLResponse('<!doctype html><html lang="ar" dir="rtl"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+escape(payload.title)+'</title><style>body{font:18px Arial;background:#eaf4fa;color:#122b41;max-width:900px;margin:auto;padding:20px}section{background:white;border-right:7px solid #13a4bd;padding:18px;margin:15px 0;border-radius:12px;break-inside:avoid}h1{color:#06647a}p{white-space:pre-wrap;line-height:1.8}@media print{body{background:white}button{display:none}}</style><button onclick="print()">🖨️ طباعة / حفظ PDF</button><h1>📘 البطاقة المرجعية — '+escape(payload.title)+'</h1><p>'+escape(payload.source)+'</p>'+blocks+'</html>',headers={"Cache-Control":"no-store"})

@router.get("/prepared")
def prepared(grade: str, subject: str, lesson: str, language: str="", format: str="pptx"):
    if format not in ("pptx","reference"):
        raise HTTPException(400,"Unsupported format")
    payload=_from_drive(grade,subject,lesson,language)
    return _pptx(payload) if format=="pptx" else _reference(payload)

@router.post("/cards")
def cards(payload: Cards, format: str="pptx"):
    if format not in ("pptx","reference"):
        raise HTTPException(400,"Unsupported format")
    return _pptx(payload) if format=="pptx" else _reference(payload)
