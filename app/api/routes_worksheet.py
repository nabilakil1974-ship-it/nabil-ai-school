"""Exports for the source-grounded interactive lesson worksheet."""

from __future__ import annotations

import base64
import io
import re
from html import escape
from pathlib import Path

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field


router = APIRouter(prefix="/worksheet", tags=["worksheet"])


class WorksheetSection(BaseModel):
    phase: str = Field(min_length=1, max_length=80)
    content: str = Field(min_length=1, max_length=40000)
    figures: list[str] = Field(default_factory=list, max_length=20)


def _figure_bytes(value: str) -> bytes | None:
    if not value.startswith("data:image/png;base64,"):
        return None
    try:
        data = base64.b64decode(value.split(",", 1)[1], validate=True)
    except Exception:
        return None
    return data if 0 < len(data) <= 5_000_000 else None


class WorksheetExport(BaseModel):
    title: str = Field(min_length=1, max_length=300)
    grade: str = Field(default="", max_length=100)
    subject: str = Field(default="", max_length=120)
    lesson: str = Field(default="", max_length=300)
    language: str = Field(default="العربية", max_length=40)
    source_label: str = Field(default="", max_length=1000)
    sections: list[WorksheetSection] = Field(min_length=1, max_length=30)


def _plain_markdown(value: str) -> str:
    text = re.sub(r"</?[A-Za-z][A-Za-z0-9]*(?:\s[^<>]*)?/?>", "", value or "")
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"[`*_]{1,3}", "", text)
    text = re.sub(r"^\s*#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"\[SOLUTION\s+(\d+)\]", r"Solution \1:", text, flags=re.IGNORECASE)
    text = re.sub(r"\[/SOLUTION\s+\d+\]", "", text, flags=re.IGNORECASE)
    return text.strip()


def _safe_filename(title: str, extension: str) -> str:
    stem = re.sub(r"[^\w\-]+", "-", title, flags=re.UNICODE).strip("-")
    return (stem[:80] or "nabil-interactive-worksheet") + extension


def _docx_bytes(request: WorksheetExport) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(12)
    normal.paragraph_format.line_spacing = 1.25

    def paragraph(text: str, *, heading: bool = False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if request.language == "العربية" else WD_ALIGN_PARAGRAPH.LEFT
        if request.language == "العربية":
            bidi = OxmlElement("w:bidi")
            bidi.set(qn("w:val"), "1")
            p._p.get_or_add_pPr().append(bidi)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(16 if heading else 12)
        run.bold = heading
        if heading:
            run.font.color.rgb = RGBColor(7, 78, 126)
        return p

    paragraph(request.title, heading=True)
    paragraph(" · ".join(x for x in (request.grade, request.subject, request.lesson) if x))
    if request.source_label:
        paragraph(request.source_label)
    for item in request.sections:
        paragraph(item.phase, heading=True)
        for line in _plain_markdown(item.content).splitlines():
            paragraph(line) if line.strip() else doc.add_paragraph()
        for encoded in item.figures:
            image = _figure_bytes(encoded)
            if image:
                doc.add_picture(io.BytesIO(image), width=Cm(16.5))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("NABIL AI — Interactive Worksheet")
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _pdf_bytes(request: WorksheetExport) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer

    font_path = next((p for p in (
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ) if p.exists()), None)
    font_name = "Helvetica"
    if font_path:
        font_name = "NabilUnicode"
        try:
            pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
        except Exception:
            pass

    def display(text: str) -> str:
        value = _plain_markdown(text)
        if request.language != "العربية":
            return value
        try:
            import arabic_reshaper
            from bidi.algorithm import get_display
            return get_display(arabic_reshaper.reshape(value))
        except Exception:
            return value

    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=1.6*cm,
        leftMargin=1.6*cm, topMargin=1.5*cm, bottomMargin=1.5*cm,
        title=request.title)
    base = getSampleStyleSheet()
    align = TA_RIGHT if request.language == "العربية" else TA_LEFT
    title_style = ParagraphStyle("NabilTitle", parent=base["Title"], fontName=font_name,
        fontSize=18, leading=24, alignment=align, textColor=colors.HexColor("#074e7e"))
    head_style = ParagraphStyle("NabilHead", parent=base["Heading2"], fontName=font_name,
        fontSize=14, leading=19, alignment=align, textColor=colors.HexColor("#0b6a8d"), spaceBefore=10)
    body_style = ParagraphStyle("NabilBody", parent=base["BodyText"], fontName=font_name,
        fontSize=11, leading=17, alignment=align, spaceAfter=5)
    def safe_paragraph(value: str, style):
        # ReportLab Paragraph interprets raw < and & as markup; worksheets
        # regularly contain inequalities such as x < 3 and x > 2.
        return Paragraph(escape(display(value)), style)

    story = [safe_paragraph(request.title, title_style)]
    meta = " · ".join(x for x in (request.grade, request.subject, request.lesson) if x)
    if meta:
        story.extend([safe_paragraph(meta, body_style), Spacer(1, 8)])
    if request.source_label:
        story.append(safe_paragraph(request.source_label, body_style))
    for item in request.sections:
        story.append(safe_paragraph(item.phase, head_style))
        for line in _plain_markdown(item.content).splitlines():
            if line.strip():
                story.append(safe_paragraph(line, body_style))
        for encoded in item.figures:
            image = _figure_bytes(encoded)
            if image:
                reader = ImageReader(io.BytesIO(image))
                width, height = reader.getSize()
                ratio = min((17.2*cm)/width, (12*cm)/height, 1)
                story.append(Image(io.BytesIO(image), width=width*ratio, height=height*ratio))
    document.build(story)
    return output.getvalue()


@router.post("/export/docx")
def export_docx(request: WorksheetExport):
    filename = _safe_filename(request.title, ".docx")
    return StreamingResponse(io.BytesIO(_docx_bytes(request)),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"})


@router.post("/export/pdf")
def export_pdf(request: WorksheetExport):
    filename = _safe_filename(request.title, ".pdf")
    return StreamingResponse(io.BytesIO(_pdf_bytes(request)), media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"})
